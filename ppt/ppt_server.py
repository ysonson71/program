import http.server
import socketserver
import io
import re
import json
import base64
import os
import google.generativeai as genai
from pptx import Presentation
from pptx.util import Inches, Pt, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

PORT = 8000
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_FILE = os.path.join(BASE_DIR, "template.pptx")
FONT_NAME = "Meiryo UI"

# 색상 상수 정의
COLOR_PRUSSIAN_BLUE = RGBColor(0, 49, 83)   # 프러시안 블루 (헤더/소제목)
COLOR_TEXT_DARK = RGBColor(30, 41, 59)       # 본문 텍스트 (#1e293b)
COLOR_MUTED = RGBColor(148, 163, 184)        # 회색 안내 문구 (#94a3b8)
COLOR_BORDER_LIGHT = RGBColor(203, 213, 225) # 연한 테두리 (#cbd5e1)

LANG_CONFIG = {
    'ja': {
        'name': '일본어',
        'expertDesc': '일본 비즈니스 기획서 및 파워포인트 원페이지 보고서 작성 전문가',
        'headerExample': '| 項目 | 現状・課題 | 推進施策 | 期待効果・KPI |',
        'toneRule': "모든 텍스트는 비즈니스 일본어 '체언지(体言止め, 명사형 종결)'로 군더더기 없이 간결하게 작성하세요.",
        'secExamples': ['### 1. 現状・課題', '### 2. 推進施策', '### 3. 期待効果・KPI'],
        'defaultTitle': 'タイトル'
    },
    'en': {
        'name': '영어',
        'expertDesc': 'Global Business Presentation & Executive Summary Specialist',
        'headerExample': '| Category | Current Status & Issues | Strategic Initiatives | Expected Impact & KPI |',
        'toneRule': 'Write all text in concise business English using noun phrases or action-oriented gerunds.',
        'secExamples': ['### 1. Current Challenges', '### 2. Key Initiatives', '### 3. Expected Outcomes'],
        'defaultTitle': 'Slide Title'
    },
    'ko': {
        'name': '한국어',
        'expertDesc': '국내 비즈니스 기획서 및 임원 보고용 원페이지 슬라이드 작성 전문가',
        'headerExample': '| 구분 | 현황 및 과제 | 추진 과제 | 기대 효과 및 KPI |',
        'toneRule': '모든 텍스트는 간결한 비즈니스 개조식(명사형 종결)으로 군더더기 없이 작성하세요.',
        'secExamples': ['### 1. 현황 및 과제', '### 2. 주요 실행 과제', '### 3. 기대 효과 및 목표'],
        'defaultTitle': '슬라이드 제목'
    }
}

def ensure_default_template():
    if not os.path.exists(TEMPLATE_FILE):
        prs = Presentation()
        prs.save(TEMPLATE_FILE)

ensure_default_template()

def set_cell_border(cell, color="000000", width="12700"):
    """셀 상하좌우 테두리 설정"""
    tcPr = cell._tc.get_or_add_tcPr()
    for border in ['lnL', 'lnR', 'lnT', 'lnB']:
        edge = parse_xml(
            f'<a:{border} {nsdecls("a")} w="{width}" cmpd="s">'
            f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
            f'</a:{border}>'
        )
        tcPr.append(edge)

def clean_text(text):
    """마크다운 태그 및 볼드 기호 정리"""
    text = re.sub(r'<br\s*/?>', '\n', text, flags=re.IGNORECASE)
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'\*(.*?)\*', r'\1', text)
    return text.replace('`', '').strip()

def ai_structure_and_translate(api_key, input_text, format_type='table', language='ja'):
    """언어(일본어 디폴트, 영어, 한국어) 및 형식에 맞추어 1줄 요약문과 본문을 구조화"""
    genai.configure(api_key=api_key)
    cfg = LANG_CONFIG.get(language, LANG_CONFIG['ja'])

    if format_type.startswith('content_image'):
        is30 = (format_type == 'content_image_30')
        guide = "(슬라이드 상측 30% 영역에는 가로로 배치될 핵심 요약 2~3개 섹션을 컴팩트하게 작성하고, 하측 70%는 대형 다이어그램/그림용으로 비워둡니다.)" if is30 else "(슬라이드 우측 50%는 그림/도표 삽입용으로 비워두며, 좌측 50%에 2~3개 섹션을 작성하세요.)"

        prompt = f"""
당신은 {cfg['expertDesc']}입니다.
입력된 내용을 분석하여, '단 1장의 슬라이드 핵심 내용' 형태로 구조화하여 {cfg['name']}로 작성하세요.
{guide}

[필수 작성 규칙]
1. 1행에는 슬라이드의 메인 타이틀을 '# 슬라이드 제목' 형식으로 작성하세요.
2. 2행에는 장표 전체를 관통하는 '핵심 내용 1줄 요약문(결론/리드문)'을 반드시 '■ 요약문' 형식으로 작성하세요.
3. 3행부터는 표(Table)를 만들지 말고, 핵심 항목(섹션)을 아래 형식으로 작성하세요:
   {cfg['secExamples'][0]}
   - 세부 내용 항목 1
   - 세부 내용 항목 2

   {cfg['secExamples'][1]}
   - 세부 내용 항목 1
   - 세부 내용 항목 2

   {cfg['secExamples'][2]}
   - 세부 내용 항목 1
   - 세부 내용 항목 2
4. {cfg['toneRule']}
5. 표(| ... |)는 절대로 출력하지 마세요.
6. '**' (볼드 기호) 또는 '<br>' (HTML 태그) 같은 기호는 사용하지 마세요.
7. 마크다운 내용 외의 다른 인사말이나 잡담은 일절 출력하지 마세요.

[입력 내용]
{input_text}
"""
    else:
        prompt = f"""
당신은 {cfg['expertDesc']}입니다.
입력된 내용을 분석하여, '단 1장의 슬라이드에 들어갈 핵심 1줄 요약문과 간결한 표(Table)' 형태로 구조화하여 {cfg['name']}로 작성하세요.

[필수 작성 규칙]
1. 1행에는 슬라이드의 메인 타이틀을 '# 슬라이드 제목' 형식으로 작성하세요.
2. 2행에는 장표 전체를 관통하는 '핵심 내용 1줄 요약문(결론/리드문)'을 반드시 '■ 요약문' 형식으로 작성하세요.
3. 그 아래에 반드시 마크다운 표(Table) 형식으로 핵심 내용을 3~4개 열, 3~5개 행으로 작성하세요.
   - 표 헤더 예시: {cfg['headerExample']}
4. 표 내부의 내용 항목에는 글머리 기호(▲, ▼, ●, •, -, * 등)를 절대로 붙이지 마세요!
5. {cfg['toneRule']}
6. '**' (볼드 기호) 또는 '<br>' (HTML 태그) 같은 기호는 사용하지 마세요.
7. 마크다운 내용 외의 다른 인사말이나 잡담은 일절 출력하지 마세요.

[입력 내용]
{input_text}
"""

    for model_name in ['gemini-2.5-flash', 'gemini-1.5-flash', 'gemini-1.5-pro']:
        try:
            model = genai.GenerativeModel(model_name)
            response = model.generate_content(prompt)
            if response and response.text:
                return response.text.strip()
        except Exception:
            continue
    raise RuntimeError("Gemini API 호출에 실패했습니다. API 키를 확인해 주세요.")

def create_base_slide(title, summary_line):
    """공통 베이스 슬라이드 생성 (16:9 와이드, 폰트: Meiryo UI, 24pt 제목, 16pt ■ 1줄 요약문)"""
    ensure_default_template()
    prs = Presentation(TEMPLATE_FILE)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    layouts = prs.slide_layouts
    layout = layouts[1] if len(layouts) > 1 else layouts[0]

    for i in range(len(prs.slides) - 1, -1, -1):
        rId = prs.slides._sldIdLst[i].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[i]

    slide = prs.slides.add_slide(layout)

    # 1. 슬라이드 제목 (24pt 볼드, Meiryo UI, 왼쪽 맞춤)
    if slide.shapes.title and slide.shapes.title.has_text_frame:
        slide.shapes.title.text = title
        slide.shapes.title.left = Inches(0.8)
        slide.shapes.title.top = Inches(0.35)
        slide.shapes.title.width = Inches(11.733)
        slide.shapes.title.height = Inches(0.65)
        if slide.shapes.title.text_frame.paragraphs:
            p_title = slide.shapes.title.text_frame.paragraphs[0]
            p_title.alignment = PP_ALIGN.LEFT
            p_title.font.name = FONT_NAME
            p_title.font.bold = True
            p_title.font.size = Pt(24)
    else:
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.733), Inches(0.65))
        p_title = tb_title.text_frame.paragraphs[0]
        p_title.text = title
        p_title.alignment = PP_ALIGN.LEFT
        p_title.font.name = FONT_NAME
        p_title.font.bold = True
        p_title.font.size = Pt(24)

    # 기본 플레이스홀더 정리
    for shape in list(slide.placeholders):
        if shape.has_text_frame and shape != slide.shapes.title:
            sp = shape._element
            sp.getparent().remove(sp)

    # 2. 핵심 1줄 요약문 (16pt 볼드, Meiryo UI, ■ 네모)
    if summary_line:
        tb_summary = slide.shapes.add_textbox(Inches(0.8), Inches(1.24), Inches(11.733), Inches(0.5))
        tf_summary = tb_summary.text_frame
        tf_summary.word_wrap = True
        p_sum = tf_summary.paragraphs[0]
        p_sum.alignment = PP_ALIGN.LEFT

        r_sq = p_sum.add_run()
        r_sq.text = "■ "
        r_sq.font.name = FONT_NAME
        r_sq.font.bold = True
        r_sq.font.size = Pt(16)
        r_sq.font.color.rgb = COLOR_TEXT_DARK

        r_txt = p_sum.add_run()
        r_txt.text = summary_line
        r_txt.font.name = FONT_NAME
        r_txt.font.bold = True
        r_txt.font.size = Pt(16)
        r_txt.font.color.rgb = COLOR_TEXT_DARK

    return prs, slide

def parse_table_content(text, language='ja'):
    """표 형식 텍스트 파싱"""
    lines = [line.strip() for line in text.strip().split('\n') if line.strip()]
    title = LANG_CONFIG.get(language, LANG_CONFIG['ja'])['defaultTitle']
    summary_line = ""
    table_rows = []

    for line in lines:
        clean_l = clean_text(line)
        if clean_l.startswith('#') and not table_rows:
            title = clean_l.lstrip('#').strip()
        elif (clean_l.startswith(('■', '□', '▶', '●', '>')) or (summary_line == "" and not clean_l.startswith('|') and not clean_l.startswith('#'))) and not table_rows:
            summary_text = re.sub(r'^[■□▶●>\s\-]+', '', clean_l).strip()
            if summary_text:
                summary_line = summary_text
        elif clean_l.startswith('|') and clean_l.endswith('|'):
            if re.match(r'^\|[\s\-:|]+\|$', clean_l):
                continue
            cells = []
            for c in clean_l.split('|')[1:-1]:
                val = clean_text(c.strip())
                val = re.sub(r'^[▲▼●○•\-\*\s]+', '', val).strip()
                cells.append(val)
            if cells:
                table_rows.append(cells)

    return title, summary_line, table_rows

def build_table_presentation(title, summary_line, table_rows):
    """1) 표 형식 슬라이드 생성 (Meiryo UI)"""
    prs, slide = create_base_slide(title, summary_line)

    if not table_rows:
        out = io.BytesIO()
        prs.save(out)
        out.seek(0)
        return out.read()

    rows = len(table_rows)
    cols = len(table_rows[0])
    left = Inches(0.8)
    top = Inches(1.85)
    width = Inches(11.733)
    height = Inches(4.8)

    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    col_width = width / cols
    for col in table.columns:
        col.width = int(col_width)

    for r_idx, row in enumerate(table_rows):
        for c_idx, val in enumerate(row):
            if c_idx >= cols:
                continue
            cell = table.cell(r_idx, c_idx)
            set_cell_border(cell, color="000000", width="12700")
            tf = cell.text_frame
            tf.word_wrap = True

            if r_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLOR_PRUSSIAN_BLUE
                p = tf.paragraphs[0]
                p.text = val
                p.font.name = FONT_NAME
                p.font.bold = True
                p.font.color.rgb = RGBColor(255, 255, 255)
                p.font.size = Pt(14)
                p.alignment = PP_ALIGN.CENTER
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(255, 255, 255)
                sub_lines = [l.strip() for l in val.split('\n') if l.strip()] or [""]
                for line_idx, line_text in enumerate(sub_lines):
                    p = tf.paragraphs[0] if line_idx == 0 else tf.add_paragraph()
                    p.alignment = PP_ALIGN.LEFT
                    p.text = re.sub(r'^[▲▼●○•\-\*\s]+', '', line_text).strip()
                    p.font.name = FONT_NAME
                    p.font.size = Pt(14)
                    p.font.color.rgb = COLOR_TEXT_DARK

    out = io.BytesIO()
    prs.save(out)
    out.seek(0)
    return out.read()

def parse_content_image(text, language='ja'):
    """2, 3) 내용+그림 형식 텍스트 파싱"""
    lines = [line.strip() for line in text.strip().split('\n') if line.strip()]
    title = LANG_CONFIG.get(language, LANG_CONFIG['ja'])['defaultTitle']
    summary_line = ""
    sections = []
    current_sec = None

    for line in lines:
        clean_l = clean_text(line)
        if clean_l.startswith('#') and not clean_l.startswith('###') and not sections:
            title = clean_l.lstrip('#').strip()
        elif (clean_l.startswith(('■', '□', '▶', '●', '>')) or (summary_line == "" and not clean_l.startswith(('#', '-')))) and not sections:
            summary_text = re.sub(r'^[■□▶●>\s\-]+', '', clean_l).strip()
            if summary_text:
                summary_line = summary_text
        elif clean_l.startswith('###') or re.match(r'^\d+[\.\)]\s*', clean_l):
            sec_title = clean_l.lstrip('#').strip()
            current_sec = {'title': sec_title, 'items': []}
            sections.append(current_sec)
        else:
            item_text = re.sub(r'^[▲▼●○•\-\*\d\.\)\s]+', '', clean_l).strip()
            if item_text:
                if current_sec is None:
                    current_sec = {'title': '', 'items': []}
                    sections.append(current_sec)
                current_sec['items'].append(item_text)

    return title, summary_line, sections

def build_content_image_presentation(title, summary_line, sections, ratio='50'):
    """2, 3) 내용+그림 슬라이드 생성
    ratio == '30': 상하 분할 (상측 30% 핵심 요약 / 하측 70% 대형 다이어그램 영역)
    ratio == '50': 좌우 분할 (좌측 50% 내용 정리 / 우측 50% 그림 영역)
    """
    prs, slide = create_base_slide(title, summary_line)

    if ratio == '30':
        # 상하 30:70 분할
        sec_count = max(len(sections), 1)
        total_w = Inches(11.733)
        col_gap = Inches(0.35) if sec_count > 1 else Inches(0)
        col_w = (total_w - (col_gap * (sec_count - 1))) / sec_count
        top_y = Inches(1.85)
        top_h = Inches(1.45)

        for sec_idx, sec in enumerate(sections):
            sec_x = Inches(0.8) + sec_idx * (col_w + col_gap)
            tb_sec = slide.shapes.add_textbox(sec_x, top_y, col_w, top_h)
            tf_sec = tb_sec.text_frame
            tf_sec.word_wrap = True
            tf_sec.margin_left = Inches(0.05)
            tf_sec.margin_right = Inches(0.05)
            tf_sec.margin_top = Inches(0.05)
            tf_sec.margin_bottom = Inches(0.05)

            is_first = True
            if sec['title']:
                p_title = tf_sec.paragraphs[0]
                is_first = False
                p_title.alignment = PP_ALIGN.LEFT
                r_title = p_title.add_run()
                r_title.text = sec['title']
                r_title.font.name = FONT_NAME
                r_title.font.bold = True
                r_title.font.size = Pt(13.5)
                r_title.font.color.rgb = COLOR_PRUSSIAN_BLUE

            for item in sec['items']:
                p_item = tf_sec.paragraphs[0] if is_first else tf_sec.add_paragraph()
                is_first = False
                p_item.alignment = PP_ALIGN.LEFT
                p_item.space_before = Pt(3)

                r_bullet = p_item.add_run()
                r_bullet.text = "• "
                r_bullet.font.name = FONT_NAME
                r_bullet.font.bold = True
                r_bullet.font.size = Pt(12)
                r_bullet.font.color.rgb = COLOR_PRUSSIAN_BLUE

                r_text = p_item.add_run()
                r_text.text = item
                r_text.font.name = FONT_NAME
                r_text.font.size = Pt(12)
                r_text.font.color.rgb = COLOR_TEXT_DARK

        # 하측 대형 그림 영역 (70%)
        bottom_y = Inches(3.55)
        bottom_h = Inches(3.40)
        guide_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), bottom_y, total_w, bottom_h)
        guide_box.fill.background()
        guide_box.line.color.rgb = COLOR_BORDER_LIGHT
        guide_box.line.width = Pt(1.5)
        try:
            guide_box.line.dash_style = MSO_LINE.DASH
        except Exception:
            pass

        tf_guide = guide_box.text_frame
        tf_guide.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_g = tf_guide.paragraphs[0]
        p_g.alignment = PP_ALIGN.CENTER
        r_g1 = p_g.add_run()
        r_g1.text = "🖼️ [ 대형 그림 / 다이어그램 삽입 영역 (70%) ]\n\n"
        r_g1.font.name = FONT_NAME
        r_g1.font.bold = True
        r_g1.font.size = Pt(14)
        r_g1.font.color.rgb = COLOR_MUTED

        r_g2 = p_g.add_run()
        r_g2.text = "(이 영역에 이미지를 자유롭게 배치하세요)"
        r_g2.font.name = FONT_NAME
        r_g2.font.size = Pt(12)
        r_g2.font.color.rgb = COLOR_MUTED

    else:
        # 좌우 50:50 분할
        left_x = Inches(0.8)
        content_y = Inches(1.85)
        content_h = Inches(5.0)
        left_w = Inches(5.67)
        right_x = Inches(6.87)
        right_w = Inches(5.67)

        tb_left = slide.shapes.add_textbox(left_x, content_y, left_w, content_h)
        tf = tb_left.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.08)
        tf.margin_right = Inches(0.08)
        tf.margin_top = Inches(0.08)
        tf.margin_bottom = Inches(0.08)

        is_first_para = True
        for sec_idx, sec in enumerate(sections):
            if sec['title']:
                p_sec = tf.paragraphs[0] if is_first_para else tf.add_paragraph()
                is_first_para = False
                p_sec.alignment = PP_ALIGN.LEFT
                if sec_idx > 0:
                    p_sec.space_before = Pt(12)

                r_sec = p_sec.add_run()
                r_sec.text = sec['title']
                r_sec.font.name = FONT_NAME
                r_sec.font.bold = True
                r_sec.font.size = Pt(15)
                r_sec.font.color.rgb = COLOR_PRUSSIAN_BLUE

            for item in sec['items']:
                p_item = tf.paragraphs[0] if is_first_para else tf.add_paragraph()
                is_first_para = False
                p_item.alignment = PP_ALIGN.LEFT
                p_item.space_before = Pt(4)

                r_bullet = p_item.add_run()
                r_bullet.text = "• "
                r_bullet.font.name = FONT_NAME
                r_bullet.font.bold = True
                r_bullet.font.size = Pt(13)
                r_bullet.font.color.rgb = COLOR_PRUSSIAN_BLUE

                r_text = p_item.add_run()
                r_text.text = item
                r_text.font.name = FONT_NAME
                r_text.font.size = Pt(13)
                r_text.font.color.rgb = COLOR_TEXT_DARK

        guide_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x, content_y, right_w, content_h)
        guide_box.fill.background()
        guide_box.line.color.rgb = COLOR_BORDER_LIGHT
        guide_box.line.width = Pt(1.5)
        try:
            guide_box.line.dash_style = MSO_LINE.DASH
        except Exception:
            pass

        tf_guide = guide_box.text_frame
        tf_guide.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_g = tf_guide.paragraphs[0]
        p_g.alignment = PP_ALIGN.CENTER
        r_g1 = p_g.add_run()
        r_g1.text = "🖼️ [ 그림 / 도표 삽입 영역 (50%) ]\n\n"
        r_g1.font.name = FONT_NAME
        r_g1.font.bold = True
        r_g1.font.size = Pt(14)
        r_g1.font.color.rgb = COLOR_MUTED

        r_g2 = p_g.add_run()
        r_g2.text = "(이 영역에 이미지를 자유롭게 배치하세요)"
        r_g2.font.name = FONT_NAME
        r_g2.font.size = Pt(12)
        r_g2.font.color.rgb = COLOR_MUTED

    out = io.BytesIO()
    prs.save(out)
    out.seek(0)
    return out.read()

class PPTHandler(http.server.BaseHTTPRequestHandler):
    def do_GET(self):
        index_path = os.path.join(BASE_DIR, "index.html")
        if os.path.exists(index_path):
            with open(index_path, "rb") as f:
                content = f.read()
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.end_headers()
            self.wfile.write(content)
        else:
            self.send_error(404, "index.html not found")

    def do_POST(self):
        content_len = int(self.headers.get('Content-Length', 0))
        post_body = self.rfile.read(content_len)
        data = json.loads(post_body.decode('utf-8'))

        if self.path == '/translate':
            try:
                api_key = data.get('api_key', '').strip()
                input_text = data.get('text', '').strip()
                format_type = data.get('format_type', 'table')
                language = data.get('language', 'ja')
                if not api_key or not input_text:
                    self.send_error(400, "API 키와 내용을 모두 입력해 주세요.")
                    return

                res_text = ai_structure_and_translate(api_key, input_text, format_type, language)
                self.send_response(200)
                self.send_header("Content-Type", "application/json; charset=utf-8")
                self.end_headers()
                self.wfile.write(json.dumps({'result': res_text}).encode('utf-8'))
            except Exception as e:
                self.send_response(500)
                self.send_header("Content-Type", "application/json; charset=utf-8")
                self.end_headers()
                self.wfile.write(json.dumps({'error': str(e)}).encode('utf-8'))

        elif self.path == '/upload_template':
            try:
                b64_tpl = data.get('template')
                if b64_tpl:
                    with open(TEMPLATE_FILE, 'wb') as f:
                        f.write(base64.b64decode(b64_tpl))
                    self.send_response(200)
                    self.end_headers()
                    self.wfile.write(b"OK")
            except Exception as e:
                self.send_error(500, str(e))

        elif self.path == '/generate':
            try:
                content_text = data.get('content', '')
                format_type = data.get('format_type', 'table')
                language = data.get('language', 'ja')

                if format_type == 'content_image_30':
                    title, summary_line, sections = parse_content_image(content_text, language)
                    out_bytes = build_content_image_presentation(title, summary_line, sections, ratio='30')
                    filename = f"presentation_content_top_image_bottom_30_70_{language}.pptx"
                elif format_type in ('content_image_50', 'content_image'):
                    title, summary_line, sections = parse_content_image(content_text, language)
                    out_bytes = build_content_image_presentation(title, summary_line, sections, ratio='50')
                    filename = f"presentation_content_image_50_50_{language}.pptx"
                else:
                    title, summary_line, table_rows = parse_table_content(content_text, language)
                    out_bytes = build_table_presentation(title, summary_line, table_rows)
                    filename = f"presentation_table_{language}.pptx"

                self.send_response(200)
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.presentationml.presentation")
                self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
                self.send_header("Content-Length", str(len(out_bytes)))
                self.end_headers()
                self.wfile.write(out_bytes)
            except Exception as e:
                self.send_response(500)
                self.send_header("Content-Type", "text/plain; charset=utf-8")
                self.end_headers()
                self.wfile.write(f"생성 오류: {str(e)}".encode('utf-8'))
        else:
            self.send_error(404)

    def log_message(self, format, *args):
        pass

if __name__ == "__main__":
    print("==================================================")
    print("AI 원페이지 파워포인트 슬라이드 생성기 실행 중")
    print(f"접속 주소: http://localhost:{PORT}")
    print(f"기본 양식 파일: {os.path.abspath(TEMPLATE_FILE)}")
    print("기본 폰트: Meiryo UI")
    print("작성 언어: 일본어(디폴트), 영어, 한국어")
    print("지원 형식: 1) 표 형식, 2) 내용+그림(50:50), 3) 내용+그림(상하 30:70)")
    print("==================================================")
    with socketserver.TCPServer(("", PORT), PPTHandler) as httpd:
        httpd.serve_forever()
